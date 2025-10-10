"""Convert PowerPoint presentations (PPTX) into structured HTML.

将 PowerPoint (PPTX) 演示文稿转换为结构化 HTML。
"""

from __future__ import annotations

import argparse
import html
import json
import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    from pptx import Presentation
    from pptx.enum.dml import MSO_FILL_TYPE, MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu
except ImportError as exc:  # pragma: no cover - defensive fallback
    raise SystemExit(
        "运行该脚本需要 python-pptx，请先执行 'pip install python-pptx' 进行安装。"
    ) from exc


def emu_to_px(value: Emu) -> float:
    """Convert PowerPoint EMU units to CSS pixels (approximation)."""
    # 1 inch = 914400 EMU and equals 96 pixels.
    return float(value) * 96.0 / 914400.0


def rgb_to_hex(rgb) -> Optional[str]:
    if rgb is None:
        return None
    try:
        return f"#{rgb:06X}"
    except (TypeError, ValueError):
        return None


def color_format_to_hex(color_format) -> Optional[str]:
    """Safely extract a hex string from a python-pptx color object."""
    if color_format is None:
        return None
    try:
        rgb = color_format.rgb  # type: ignore[attr-defined]
    except AttributeError:
        return None
    return rgb_to_hex(rgb)


def serialize_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, str)):
        return value
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:  # pragma: no cover - defensive
            pass
    return str(value)


def alignment_to_css(alignment: Optional[PP_ALIGN]) -> str:
    if alignment is None:
        return ""
    mapping = {
        PP_ALIGN.LEFT: "left",
        PP_ALIGN.CENTER: "center",
        PP_ALIGN.RIGHT: "right",
        PP_ALIGN.JUSTIFY: "justify",
    }
    value = mapping.get(alignment)
    if value:
        return f"text-align: {value};"
    return ""


def build_text_html(shape) -> str:
    text_frame = shape.text_frame
    paragraphs_html: List[str] = []
    for paragraph in text_frame.paragraphs:
        runs_html: List[str] = []
        para_styles: List[str] = []

        if paragraph.space_before is not None:
            para_styles.append(f"margin-top: {paragraph.space_before.pt:.2f}pt;")
        if paragraph.space_after is not None:
            para_styles.append(f"margin-bottom: {paragraph.space_after.pt:.2f}pt;")
        alignment_style = alignment_to_css(paragraph.alignment)
        if alignment_style:
            para_styles.append(alignment_style)

        for run in paragraph.runs:
            styles: List[str] = []
            font = run.font
            if font.bold:
                styles.append("font-weight: bold;")
            if font.italic:
                styles.append("font-style: italic;")
            if font.underline:
                styles.append("text-decoration: underline;")
            if font.size:
                styles.append(f"font-size: {font.size.pt:.2f}pt;")
            if font.name:
                styles.append(f"font-family: '{font.name}';")
            color_hex = color_format_to_hex(font.color)
            if color_hex:
                styles.append(f"color: {color_hex};")
            run_html = html.escape(run.text)
            if not run_html:
                continue
            if styles:
                runs_html.append(f"<span style=\"{''.join(styles)}\">{run_html}</span>")
            else:
                runs_html.append(run_html)
        paragraph_style_attr = ''.join(para_styles)
        level = paragraph.level if paragraph.level is not None else 0
        level_attr = f" data-level=\"{level}\""
        paragraphs_html.append(
            f"<p{level_attr} style=\"{paragraph_style_attr}\">{''.join(runs_html)}</p>"
        )
    return ''.join(paragraphs_html)


def extract_picture(shape, assets_dir: Path, slide_idx: int, shape_idx: int) -> str:
    image = shape.image
    ext = image.ext
    filename = f"slide{slide_idx:03d}_shape{shape_idx:03d}{ext}"
    output_path = assets_dir / filename
    with output_path.open('wb') as fp:
        fp.write(image.blob)
    return f"<img src=\"{assets_dir.name}/{filename}\" alt=\"{html.escape(shape.name)}\" style=\"width: 100%; height: 100%; object-fit: contain;\">"


def build_table_html(shape) -> str:
    rows_html: List[str] = []
    table = shape.table
    for row in table.rows:
        cells_html: List[str] = []
        for cell in row.cells:
            cell_text = html.escape(cell.text)
            cell_styles: List[str] = []
            color = color_format_to_hex(getattr(cell.fill, "fore_color", None) if cell.fill else None)
            if color:
                cell_styles.append(f"background-color: {color};")
            cells_html.append(f"<td style=\"{''.join(cell_styles)}\">{cell_text}</td>")
        rows_html.append(f"<tr>{''.join(cells_html)}</tr>")
    return f"<table>{''.join(rows_html)}</table>"


def chart_to_html(chart) -> Tuple[str, Dict[str, Any]]:
    headers = ["Category"]
    series_meta: List[Dict[str, Any]] = []
    for idx, series in enumerate(chart.series):
        name = series.name if series.name is not None else f"Series {idx + 1}"
        headers.append(str(name))
        series_meta.append(
            {
                "name": None if series.name is None else str(series.name),
                "values": [serialize_value(value) for value in (series.values or [])],
            }
        )

    plots = list(chart.plots)
    categories = []
    if plots and getattr(plots[0], "has_categories", False):
        try:
            categories = list(plots[0].categories)
        except TypeError:
            categories = []
    if not categories:
        max_len = max((len(series.values) for series in chart.series), default=0)
        categories = [f"Point {idx + 1}" for idx in range(max_len)]

    header_html = "<tr>" + ''.join(f"<th>{html.escape(str(item))}</th>" for item in headers) + "</tr>"

    rows_html: List[str] = []
    categories_text: List[str] = []
    for cat_idx, category in enumerate(categories):
        category_label = getattr(category, "label", category)
        if isinstance(category_label, (list, tuple)):
            category_text = " / ".join(str(item) for item in category_label)
        else:
            category_text = str(category_label)
        categories_text.append(category_text)
        row_cells = [html.escape(category_text)]
        for series in chart.series:
            values = series.values or []
            value = values[cat_idx] if cat_idx < len(values) else ""
            row_cells.append(html.escape("" if value is None else str(serialize_value(value))))
        rows_html.append("<tr>" + ''.join(f"<td>{cell}</td>" for cell in row_cells) + "</tr>")

    chart_meta: Dict[str, Any] = {
        "type": str(chart.chart_type),
        "has_legend": getattr(chart, "has_legend", False),
        "series": series_meta,
        "categories": [serialize_value(text) for text in categories_text],
    }
    if getattr(chart, "has_title", False) and chart.chart_title is not None:
        chart_meta["title"] = chart.chart_title.text_frame.text

    html_table = f"<table class=\"chart-data\">{header_html}{''.join(rows_html)}</table>"
    return html_table, chart_meta


def shape_to_html(shape, assets_dir: Path, slide_idx: int, shape_idx: int) -> str:
    base_styles = [
        "position: absolute;",
        f"left: {emu_to_px(shape.left):.2f}px;",
        f"top: {emu_to_px(shape.top):.2f}px;",
        f"width: {emu_to_px(shape.width):.2f}px;",
        f"height: {emu_to_px(shape.height):.2f}px;",
    ]
    rotation = getattr(shape, "rotation", 0) or 0
    if rotation:
        base_styles.append(f"transform: rotate({rotation:.2f}deg);")

    background_color = None
    fill = getattr(shape, "fill", None)
    if fill is not None and getattr(fill, "type", None) == MSO_FILL_TYPE.SOLID:
        background_color = color_format_to_hex(getattr(fill, "fore_color", None))
    if background_color:
        base_styles.append(f"background-color: {background_color};")

    line = getattr(shape, "line", None)
    line_fill = getattr(line, "fill", None) if line is not None else None
    dash_style = getattr(line, "dash_style", None) if line is not None else None
    border_color = None
    if line_fill is not None and getattr(line_fill, "type", None) == MSO_FILL_TYPE.SOLID:
        border_color = color_format_to_hex(getattr(line_fill, "fore_color", None))
    if border_color:
        width = emu_to_px(line.width) if getattr(line, "width", None) else 1.0
        dash_mapping = {
            MSO_LINE_DASH_STYLE.DASH: "dashed",
            MSO_LINE_DASH_STYLE.DOT: "dotted",
            MSO_LINE_DASH_STYLE.DASH_DOT: "dash-dot",
            MSO_LINE_DASH_STYLE.DASH_DOT_DOT: "dash-dot-dot",
            MSO_LINE_DASH_STYLE.LONG_DASH: "dashed",
            MSO_LINE_DASH_STYLE.LONG_DASH_DOT: "dash-dot",
            MSO_LINE_DASH_STYLE.SOLID: "solid",
            MSO_LINE_DASH_STYLE.ROUND_DOT: "dotted",
            MSO_LINE_DASH_STYLE.SYSTEM_DASH: "dashed",
            MSO_LINE_DASH_STYLE.SYSTEM_DASH_DOT: "dash-dot",
            MSO_LINE_DASH_STYLE.SYSTEM_DASH_DOT_DOT: "dash-dot-dot",
            MSO_LINE_DASH_STYLE.SYSTEM_DOT: "dotted",
        }
        css_dash = dash_mapping.get(dash_style, "solid")
        base_styles.append(f"border: {width:.2f}px {css_dash} {border_color};")

    data_attrs = [
        f"data-shape-id=\"{shape.shape_id}\"",
        f"data-shape-name=\"{html.escape(shape.name)}\"",
        f"data-shape-type=\"{shape.shape_type}\"",
    ]
    if getattr(shape, "is_placeholder", False):
        placeholder_type = getattr(getattr(shape, "placeholder_format", None), "type", None)
        if placeholder_type is not None:
            data_attrs.append(f"data-placeholder-type=\"{placeholder_type}\"")

    metadata: Dict[str, Any] = {
        "shape_id": shape.shape_id,
        "name": shape.name,
        "shape_type": str(shape.shape_type),
        "left_emu": int(shape.left),
        "top_emu": int(shape.top),
        "width_emu": int(shape.width),
        "height_emu": int(shape.height),
        "rotation": float(rotation),
        "z_order": getattr(shape, "z_order_position", None),
        "raw_xml": shape.element.xml,
    }
    if background_color:
        metadata["fill_color"] = background_color
    if border_color:
        metadata["line_color"] = border_color
        metadata["line_width_emu"] = int(line.width) if getattr(line, "width", None) else None
        metadata["line_dash_style"] = str(dash_style) if dash_style is not None else None
    if getattr(shape, "is_placeholder", False):
        placeholder_type = getattr(getattr(shape, "placeholder_format", None), "type", None)
        if placeholder_type is not None:
            metadata["placeholder_type"] = str(placeholder_type)

    inner_html = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        inner_html = extract_picture(shape, assets_dir, slide_idx, shape_idx)
        image_obj = shape.image
        dpi = getattr(image_obj, "dpi", None)
        metadata["image"] = {
            "content_type": getattr(image_obj, "content_type", None),
            "ext": getattr(image_obj, "ext", None),
            "dpi": {
                "horz": dpi[0] if dpi else None,
                "vert": dpi[1] if dpi else None,
            },
        }
    elif shape.has_text_frame:
        inner_html = build_text_html(shape)
        metadata["text"] = shape.text
    elif shape.has_table:
        inner_html = build_table_html(shape)
        metadata["table"] = {
            "rows": len(shape.table.rows),
            "columns": len(shape.table.columns),
        }
    elif getattr(shape, "has_chart", False):
        inner_html, chart_meta = chart_to_html(shape.chart)
        data_attrs.append(f"data-chart-type=\"{shape.chart.chart_type}\"")
        metadata["chart"] = chart_meta
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children_html: List[str] = []
        for child_idx, child in enumerate(shape.shapes):
            children_html.append(shape_to_html(child, assets_dir, slide_idx, child_idx))
        inner_html = ''.join(children_html)
        metadata["group"] = {"count": len(shape.shapes)}
    else:
        inner_html = f"<pre>{html.escape(str(shape.text if hasattr(shape, 'text') else ''))}</pre>"
        if hasattr(shape, "text"):
            metadata["text"] = shape.text

    style_attr = ''.join(base_styles)
    metadata_json = html.escape(json.dumps(metadata, ensure_ascii=False, separators=(",", ":")))
    metadata_tag = (
        f"<script type=\"application/json\" class=\"shape-metadata\">{metadata_json}</script>"
    )
    return f"<div class=\"shape\" {' '.join(data_attrs)} style=\"{style_attr}\">{inner_html}{metadata_tag}</div>"


def slide_background_style(slide, assets_dir: Path, slide_idx: int) -> str:
    fill = slide.background.fill
    if not fill:
        return ""
    try:
        if fill.type == MSO_FILL_TYPE.SOLID:
            color = color_format_to_hex(getattr(fill, "fore_color", None))
            if color:
                return f"background-color: {color};"
        elif fill.type == MSO_FILL_TYPE.PICTURE:
            image = fill.picture.image
            if image is not None:
                filename = f"slide{slide_idx:03d}_background{image.ext}"
                output_path = assets_dir / filename
                if not output_path.exists():
                    with output_path.open("wb") as fp:
                        fp.write(image.blob)
                return (
                    f"background-image: url('{assets_dir.name}/{filename}'); "
                    "background-size: cover; background-position: center;"
                )
    except AttributeError:
        return ""
    return ""


def slide_to_html(slide, assets_dir: Path, slide_idx: int) -> str:
    shapes_html: List[str] = []
    for shape_idx, shape in enumerate(slide.shapes):
        shapes_html.append(shape_to_html(shape, assets_dir, slide_idx, shape_idx))

    background_style = slide_background_style(slide, assets_dir, slide_idx)
    title_text = slide.shapes.title.text if slide.shapes.title else f"Slide {slide_idx + 1}"
    title = html.escape(title_text)
    notes_text = ""
    if slide.has_notes_slide:
        notes_text = html.escape(slide.notes_slide.notes_text_frame.text)
    layout_name = html.escape(slide.slide_layout.name) if slide.slide_layout else ""
    slide_metadata: Dict[str, Any] = {
        "index": slide_idx,
        "slide_id": slide.slide_id,
        "title": title_text,
        "layout": slide.slide_layout.name if slide.slide_layout else None,
        "notes": slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else "",
        "shape_count": len(slide.shapes),
        "raw_xml": slide.element.xml,
    }
    metadata_json = html.escape(json.dumps(slide_metadata, ensure_ascii=False, separators=(",", ":")))
    metadata_tag = (
        f"<script type=\"application/json\" class=\"slide-metadata\">{metadata_json}</script>"
    )
    return (
        f"<section class=\"slide\" data-title=\"{title}\" data-layout=\"{layout_name}\" style=\"{background_style}\">"
        f"{''.join(shapes_html)}"
        f"<aside class=\"notes\">{notes_text}</aside>"
        f"{metadata_tag}"
        "</section>"
    )


def generate_html(presentation: Presentation, assets_dir: Path) -> str:
    slides_html: List[str] = []
    for slide_idx, slide in enumerate(presentation.slides):
        slides_html.append(slide_to_html(slide, assets_dir, slide_idx))

    core_props = {}
    props = presentation.core_properties
    for attr in [
        "author",
        "category",
        "comments",
        "content_status",
        "created",
        "identifier",
        "keywords",
        "language",
        "last_modified_by",
        "last_printed",
        "modified",
        "revision",
        "subject",
        "title",
        "version",
    ]:
        value = getattr(props, attr, None)
        if value not in (None, ""):
            core_props[attr] = serialize_value(value)

    custom_props = {}
    try:
        custom_properties = presentation.custom_properties
        for name in custom_properties:
            custom_props[name] = serialize_value(custom_properties[name])
    except AttributeError:
        custom_props = {}

    presentation_metadata = {
        "slide_count": len(presentation.slides),
        "core_properties": core_props,
        "custom_properties": custom_props,
    }
    metadata_json = html.escape(
        json.dumps(presentation_metadata, ensure_ascii=False, separators=(",", ":"))
    )
    presentation_metadata_tag = (
        f"<script type=\"application/json\" id=\"presentation-metadata\">{metadata_json}</script>"
    )

    css = """
    body { font-family: Arial, Helvetica, sans-serif; margin: 0; padding: 0; background: #111; }
    .deck { display: flex; flex-direction: column; gap: 48px; padding: 32px; }
    .slide { position: relative; width: 960px; height: 540px; margin: 0 auto; background: white; overflow: hidden; box-shadow: 0 4px 24px rgba(0, 0, 0, 0.25); }
    .slide .shape { box-sizing: border-box; }
    .slide .notes { display: none; }
    table { border-collapse: collapse; }
    table.chart-data { width: 100%; margin-top: 8px; }
    table.chart-data th, table.chart-data td { border: 1px solid #cccccc; padding: 4px 8px; }
    """

    return (
        "<!DOCTYPE html>"
        "<html lang=\"en\">"
        "<head>"
        "<meta charset=\"utf-8\">"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">"
        "<title>PPTX to HTML Export</title>"
        f"<style>{css}</style>"
        "</head>"
        "<body>"
        f"{presentation_metadata_tag}"
        "<main class=\"deck\">"
        f"{''.join(slides_html)}"
        "</main>"
        "</body>"
        "</html>"
    )


def convert_pptx_to_html(input_path: Path, output_html: Path) -> None:
    presentation = Presentation(str(input_path))

    assets_dir = output_html.with_name(output_html.stem + "_assets")
    if assets_dir.exists():
        shutil.rmtree(assets_dir)
    assets_dir.mkdir(parents=True, exist_ok=True)

    html_content = generate_html(presentation, assets_dir)
    output_html.write_text(html_content, encoding="utf-8")



def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert PPTX presentations to HTML.")
    parser.add_argument("input", type=Path, help="Path to the PPTX file to convert.")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Path for the output HTML file. Defaults to <input_name>.html",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_path: Path = args.input
    if not input_path.exists():
        raise SystemExit(f"Input file '{input_path}' does not exist.")
    if input_path.suffix.lower() != ".pptx":
        raise SystemExit("Input file must have a .pptx extension.")

    output_html: Path
    if args.output is None:
        output_html = input_path.with_suffix(".html")
    else:
        output_html = args.output
        if output_html.is_dir():
            output_html = output_html / (input_path.stem + ".html")
        if output_html.suffix.lower() != ".html":
            output_html = output_html.with_suffix(".html")

    convert_pptx_to_html(input_path, output_html)
    print(f"Exported '{input_path}' to '{output_html}' with assets in '{output_html.stem}_assets/'.")


if __name__ == "__main__":
    main()
